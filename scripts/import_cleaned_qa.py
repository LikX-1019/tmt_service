import argparse
import json
import re
from datetime import datetime
from pathlib import Path
from typing import Any

import openpyxl

REQUIRED_HEADERS = {
    "问题编号",
    "商品编码",
    "商品名称",
    "售前售后",
    "问题类型",
    "标准问法",
    "同义问法/触发表达",
    "检索关键词",
    "标准回答",
    "风险等级",
    "必答要点",
    "禁止表达",
    "适用资料版本",
    "当前话术版本",
    "QA建议状态",
}

SERVICE_STAGE_MAP = {"售前": "pre_sale", "售后": "post_sale", "通用": "general"}
RISK_LEVEL_MAP = {"低": "low", "中": "medium", "高": "high"}
REVIEW_STATUS_MAP = {
    "可用": ("published", "usable"),
    "待复核": ("draft", "pending_review"),
    "待验证": ("draft", "pending_validation"),
}

COLUMNS = [
    "qa_code",
    "product_code",
    "product_name",
    "standard_question",
    "standard_answer",
    "similar_questions",
    "keywords",
    "intent_code",
    "question_type",
    "service_stage",
    "required_points",
    "prohibited_expressions",
    "risk_level",
    "need_human",
    "applicable_version",
    "answer_version",
    "priority",
    "status",
    "review_status",
    "source",
    "import_batch_no",
    "source_row_no",
    "effective_at",
    "expired_at",
    "created_by",
    "approved_by",
]


def text(value: Any) -> str | None:
    if value is None:
        return None
    normalized = str(value).strip()
    return normalized or None


def split_items(value: Any, *, split_whitespace: bool = False) -> list[str]:
    raw = text(value)
    if not raw:
        return []
    pattern = r"[\s；;]+" if split_whitespace else r"[\r\n；;]+"
    result = []
    seen = set()
    for item in re.split(pattern, raw):
        item = item.strip()
        if item and item not in seen:
            seen.add(item)
            result.append(item)
    return result


def answer_version(version_code: Any) -> int:
    match = re.search(r"-V(\d+)$", text(version_code) or "", flags=re.IGNORECASE)
    return int(match.group(1)) if match else 1


def sql_value(value: Any) -> str:
    if value is None:
        return "NULL"
    if isinstance(value, bool):
        return "1" if value else "0"
    if isinstance(value, int):
        return str(value)
    encoded = str(value).encode("utf-8").hex()
    return f"CONVERT(0x{encoded} USING utf8mb4)"


def load_rows(source: Path, sheet_name: str, batch_no: str) -> list[dict[str, Any]]:
    workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"缺少工作表：{sheet_name}")
    sheet = workbook[sheet_name]
    iterator = sheet.iter_rows(min_row=5, values_only=True)
    headers = list(next(iterator))
    missing_headers = sorted(REQUIRED_HEADERS - set(headers))
    if missing_headers:
        raise ValueError(f"缺少字段：{', '.join(missing_headers)}")

    imported: list[dict[str, Any]] = []
    seen_codes: set[str] = set()
    errors: list[str] = []
    for source_row_no, values in enumerate(iterator, start=6):
        if not values[0]:
            continue
        row = dict(zip(headers, values))
        qa_code = text(row.get("问题编号"))
        question = text(row.get("标准问法"))
        answer = text(row.get("标准回答"))
        if not qa_code or not question or not answer:
            errors.append(f"第 {source_row_no} 行缺少问题编号、标准问法或标准回答")
            continue
        if qa_code in seen_codes:
            errors.append(f"第 {source_row_no} 行问题编号重复：{qa_code}")
            continue
        seen_codes.add(qa_code)

        stage = SERVICE_STAGE_MAP.get(text(row.get("售前售后")) or "")
        risk = RISK_LEVEL_MAP.get(text(row.get("风险等级")) or "")
        publish_and_review = REVIEW_STATUS_MAP.get(text(row.get("QA建议状态")) or "")
        if not stage:
            errors.append(f"第 {source_row_no} 行售前售后值不支持：{row.get('售前售后')}")
        if not risk:
            errors.append(f"第 {source_row_no} 行风险等级值不支持：{row.get('风险等级')}")
        if not publish_and_review:
            errors.append(f"第 {source_row_no} 行 QA 建议状态不支持：{row.get('QA建议状态')}")
        if not stage or not risk or not publish_and_review:
            continue

        status, review_status = publish_and_review
        aliases = [item for item in split_items(row.get("同义问法/触发表达")) if item != question]
        product_code = text(row.get("商品编码"))
        if product_code in {"通用", "COMMON"}:
            product_code = None

        imported.append(
            {
                "qa_code": qa_code,
                "product_code": product_code,
                "product_name": text(row.get("商品名称")),
                "standard_question": question,
                "standard_answer": answer,
                "similar_questions": json.dumps(aliases, ensure_ascii=False),
                "keywords": json.dumps(
                    split_items(row.get("检索关键词"), split_whitespace=True),
                    ensure_ascii=False,
                ),
                "intent_code": None,
                "question_type": text(row.get("问题类型")),
                "service_stage": stage,
                "required_points": json.dumps(
                    split_items(row.get("必答要点")), ensure_ascii=False
                ),
                "prohibited_expressions": json.dumps(
                    split_items(row.get("禁止表达")), ensure_ascii=False
                ),
                "risk_level": risk,
                "need_human": 0,
                "applicable_version": text(row.get("适用资料版本")),
                "answer_version": answer_version(row.get("当前话术版本")),
                "priority": 0,
                "status": status,
                "review_status": review_status,
                "source": source.name,
                "import_batch_no": batch_no,
                "source_row_no": source_row_no,
                "effective_at": None,
                "expired_at": None,
                "created_by": "data_import",
                "approved_by": None,
            }
        )

    if errors:
        raise ValueError("\n".join(errors))
    return imported


def render_sql(rows: list[dict[str, Any]], batch_size: int = 50) -> str:
    update_columns = [column for column in COLUMNS if column != "qa_code"]
    statements = ["SET NAMES utf8mb4;", "START TRANSACTION;"]
    quoted_columns = ", ".join(f"`{column}`" for column in COLUMNS)
    updates = ", ".join(f"`{column}` = VALUES(`{column}`)" for column in update_columns)
    for start in range(0, len(rows), batch_size):
        batch = rows[start : start + batch_size]
        values = []
        for row in batch:
            values.append("(" + ", ".join(sql_value(row[column]) for column in COLUMNS) + ")")
        statements.append(
            f"INSERT INTO `cs_qa` ({quoted_columns}) VALUES\n"
            + ",\n".join(values)
            + f"\nON DUPLICATE KEY UPDATE {updates};"
        )
    statements.append("COMMIT;")
    return "\n\n".join(statements) + "\n"


def main() -> None:
    parser = argparse.ArgumentParser(description="Convert cleaned QA workbook to MySQL SQL")
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--sheet", default="标准QA")
    parser.add_argument("--batch-no", default=datetime.now().strftime("QA-%Y%m%d-%H%M%S"))
    arguments = parser.parse_args()

    rows = load_rows(arguments.source, arguments.sheet, arguments.batch_no)
    if not rows:
        raise ValueError("没有可导入的 QA 数据")
    arguments.output.parent.mkdir(parents=True, exist_ok=True)
    arguments.output.write_text(render_sql(rows), encoding="utf-8")
    summary = {
        "batch_no": arguments.batch_no,
        "rows": len(rows),
        "published": sum(row["status"] == "published" for row in rows),
        "draft": sum(row["status"] == "draft" for row in rows),
        "usable": sum(row["review_status"] == "usable" for row in rows),
        "pending_review": sum(row["review_status"] == "pending_review" for row in rows),
        "pending_validation": sum(
            row["review_status"] == "pending_validation" for row in rows
        ),
        "distinct_alias_rows": sum(row["similar_questions"] != "[]" for row in rows),
    }
    print(json.dumps(summary, ensure_ascii=False))


if __name__ == "__main__":
    main()
