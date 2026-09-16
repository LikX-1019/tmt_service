"""用逐条标注样本校准 RAG 自动发送阈值。"""

from __future__ import annotations

import argparse
import asyncio
import csv
import json
from dataclasses import asdict
from datetime import datetime, timezone
from hashlib import sha256
from pathlib import Path
from typing import Any

from app.api.dependencies import get_qa_service
from app.core.config import get_settings
from app.services.auto_reply_policy import choose_calibration


REQUIRED_COLUMNS = {"query", "expected_qa"}


def _normalize_rows(rows: list[dict[str, Any]]) -> list[dict[str, str]]:
    result: list[dict[str, str]] = []
    for row in rows:
        query = str(row.get("query") or "").strip()
        expected = str(row.get("expected_qa") or "").strip()
        if query and expected:
            result.append({"query": query, "expected_qa": expected})
    return result


def _load_labels(path: Path, sheet_name: str) -> list[dict[str, str]]:
    suffix = path.suffix.lower()
    if suffix == ".jsonl":
        rows = [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]
    elif suffix == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as handle:
            rows = list(csv.DictReader(handle))
    elif suffix == ".xlsx":
        import openpyxl

        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"标注工作簿缺少工作表：{sheet_name}")
        values = workbook[sheet_name].iter_rows(values_only=True)
        headers = [str(value or "").strip() for value in next(values)]
        rows = [dict(zip(headers, row)) for row in values]
    else:
        raise ValueError("标注文件仅支持 .xlsx、.csv 或 .jsonl")
    if rows and not REQUIRED_COLUMNS.issubset(rows[0]):
        raise ValueError("标注文件必须包含 query 和 expected_qa 两列")
    return _normalize_rows(rows)


async def calibrate(labels_path: Path, sheet_name: str, limit: int | None) -> dict[str, Any]:
    settings = get_settings()
    labels = await asyncio.to_thread(_load_labels, labels_path, sheet_name)
    if limit:
        labels = labels[:limit]
    service = await get_qa_service()
    records: list[dict[str, Any]] = []
    for index, row in enumerate(labels, 1):
        candidates = await service.retrieve_rag_candidates(row["query"])
        top = candidates[0] if candidates else None
        second = candidates[1] if len(candidates) > 1 else None
        top_score = top.rerank_score if top else None
        margin = None
        if top_score is not None:
            margin = top_score - second.rerank_score if second and second.rerank_score is not None else top_score
        records.append(
            {
                "expected_qa": row["expected_qa"],
                "predicted_qa": top.chunk_id if top else None,
                "top_score": top_score,
                "margin": margin,
            }
        )
        if index % 20 == 0:
            print(f"已评估 {index}/{len(labels)} 条")
    selected = choose_calibration(
        records,
        min_precision=settings.auto_reply_min_precision,
        min_samples=settings.auto_reply_min_samples,
    )
    payload = {
        **asdict(selected),
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "labels_sha256": sha256(labels_path.read_bytes()).hexdigest(),
        "labels_file": labels_path.name,
        "policy_version": settings.auto_reply_policy_version,
    }
    output = settings.auto_reply_calibration_path
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return payload


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("labels", type=Path, help="逐条标注文件，列名为 query、expected_qa")
    parser.add_argument("--sheet", default="质检标注", help="XLSX 工作表名称")
    parser.add_argument("--limit", type=int, default=None)
    args = parser.parse_args()
    if not args.labels.is_file():
        raise SystemExit(f"标注文件不存在：{args.labels}")
    result = asyncio.run(calibrate(args.labels, args.sheet, args.limit))
    print(json.dumps(result, ensure_ascii=False, indent=2))
    if not result["enabled"]:
        print("RAG 自动回复保持关闭：样本量或接受集合精确率未达门槛。")


if __name__ == "__main__":
    main()
