from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import date, datetime, time, timezone
from decimal import Decimal
from io import BytesIO
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from pydantic import ValidationError

from app.schemas import ProductCreate, VariantCreate


PRODUCT_SHEET = "商品导入"
VARIANT_SHEET = "SKU导入"
EXAMPLE_SHEET = "填写示例"
HELP_SHEET = "字段说明"
MAX_IMPORT_ROWS = 1000

PRODUCT_COLUMNS = [
    ("平台商品ID*", "id", "必填；平台商品卡片 goods_id，最长 64 字", "972793561880"),
    ("内部商品编码", "internal_code", "选填；公司内部唯一商品编码", "BG07"),
    ("平台", "platform", "选填；拼多多、京东、淘宝/天猫、其他，默认拼多多", "拼多多"),
    ("资料状态", "status", "选填；草稿、已发布、已下线，默认草稿", "已发布"),
    ("商品名称*", "name", "必填；消费者看到的完整商品名称", "JAFFICK 运动护腕支架 Pro"),
    ("品牌", "brand", "选填；包装或平台公示品牌", "JAFFICK"),
    ("型号", "model", "选填；商品型号", "BG07 Pro"),
    ("类目", "category_name", "选填；建议使用“一级 / 二级 / 三级”", "运动户外 / 运动护具 / 护腕"),
    ("商品简介*", "summary", "必填；1～3 句可核实事实，最长 4000 字", "用于日常运动及办公场景下的手腕支撑。"),
    ("核心卖点", "selling_points", "选填；多个卖点用换行或 | 分隔，最多 30 条", "三段式可调节绑带|左右手均可佩戴|透气面料"),
    ("商品规格", "specifications", "选填；每项写“名称=值”，多个用换行或 | 分隔", "颜色=黑色|尺码=均码|材质=锦纶、聚酯纤维、氨纶"),
    ("使用方法", "usage", "选填；安装、佩戴、清洁等操作顺序", "松开绑带，调整至舒适松紧度后粘合。"),
    ("适用人群/场景", "suitable_for", "选填；适用对象和使用场景", "适合日常运动、办公及一般手腕支撑需求。"),
    ("注意事项", "warnings", "选填；禁忌、异常处理及风险提示", "出现麻木、红肿或疼痛加重时立即停止使用。"),
    ("售后限制", "after_sales_limits", "选填；拆封、使用、污染等退换边界", "已使用或污染的贴身类商品不支持无理由退换。"),
    ("资料负责人", "owner_name", "选填；负责维护资料的运营人员", "张琳"),
    ("审核人", "reviewed_by", "选填；审核商品事实和合规口径的人员", "李明"),
    ("生效时间", "effective_at", "选填；格式 2026-09-18 10:30", "2026-09-18 10:30"),
    ("发布时间", "published_at", "选填；格式 2026-09-18 10:30", "2026-09-18 10:30"),
    ("合规备注", "compliance_notes", "选填；资质依据、禁用表述等", "不得描述为治疗产品，不得承诺治愈。"),
]

VARIANT_COLUMNS = [
    ("平台商品ID*", "product_id", "必填；必须对应商品导入表或数据库中已有商品", "972793561880"),
    ("SKU ID*", "sku_id", "必填；稳定且唯一的 SKU 编码", "BG07-BLACK-M"),
    ("SKU名称*", "name", "必填；顾客容易理解的规格组合名称", "黑色 / M 码"),
    ("规格属性", "attributes", "选填；每项写“名称=值”，多个用换行或 | 分隔", "颜色=黑色|尺码=M"),
    ("币种", "currency", "选填；三位币种代码，默认 CNY", "CNY"),
    ("价格", "price", "选填；非负数字，最多两位小数", 69.90),
    ("库存状态", "stock_status", "选填；未知、有货、库存紧张、缺货", "有货"),
    ("库存数量", "stock_quantity", "选填；非负整数", 25),
    ("SKU状态", "status", "选填；启用、停用，默认启用", "启用"),
]

PLATFORM_MAP = {"拼多多": "pdd", "pdd": "pdd", "京东": "jd", "jd": "jd", "淘宝": "taobao", "天猫": "taobao", "淘宝/天猫": "taobao", "taobao": "taobao", "其他": "other", "other": "other"}
PRODUCT_STATUS_MAP = {"草稿": "draft", "draft": "draft", "已发布": "published", "published": "published", "已下线": "offline", "offline": "offline"}
STOCK_STATUS_MAP = {"未知": "unknown", "unknown": "unknown", "有货": "in_stock", "in_stock": "in_stock", "库存紧张": "low_stock", "low_stock": "low_stock", "缺货": "out_of_stock", "out_of_stock": "out_of_stock"}
VARIANT_STATUS_MAP = {"启用": "active", "active": "active", "停用": "inactive", "inactive": "inactive"}


@dataclass(slots=True)
class ImportErrorItem:
    sheet: str
    row: int
    message: str


@dataclass(slots=True)
class VariantImport:
    row: int
    product_id: str
    variant: VariantCreate


@dataclass(slots=True)
class ParsedCatalog:
    products: list[ProductCreate] = field(default_factory=list)
    variants: list[VariantImport] = field(default_factory=list)
    errors: list[ImportErrorItem] = field(default_factory=list)


def _style_sheet(sheet, columns: list[tuple[str, str, str, Any]], *, example: bool = False) -> None:
    header_fill = PatternFill("solid", fgColor="126B57")
    required_fill = PatternFill("solid", fgColor="0E5948")
    sample_fill = PatternFill("solid", fgColor="FFF4D6")
    border = Border(bottom=Side(style="thin", color="D9E0E3"))
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}1"
    sheet.row_dimensions[1].height = 30
    for index, (title, _key, description, sample) in enumerate(columns, start=1):
        cell = sheet.cell(1, index, title)
        cell.fill = required_fill if title.endswith("*") else header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.comment = None
        width = max(14, min(36, max(len(title) * 2, len(description) // 2)))
        sheet.column_dimensions[get_column_letter(index)].width = width
        if example:
            sample_cell = sheet.cell(2, index, sample)
            sample_cell.fill = sample_fill
            sample_cell.border = border
            sample_cell.alignment = Alignment(vertical="top", wrap_text=True)
    if example:
        sheet.row_dimensions[2].height = 76


def _add_validations(sheet, product: bool) -> None:
    validations = (
        [("C", '"拼多多,京东,淘宝/天猫,其他"'), ("D", '"草稿,已发布,已下线"')]
        if product
        else [("G", '"未知,有货,库存紧张,缺货"'), ("I", '"启用,停用"')]
    )
    for column, formula in validations:
        validation = DataValidation(type="list", formula1=formula, allow_blank=True)
        validation.error = "请选择下拉列表中的值"
        validation.errorTitle = "填写值无效"
        sheet.add_data_validation(validation)
        validation.add(f"{column}2:{column}{MAX_IMPORT_ROWS + 1}")
    required_columns = [1, 5, 9] if product else [1, 2, 3]
    red_fill = PatternFill("solid", fgColor="FDE9E7")
    for column in required_columns:
        letter = get_column_letter(column)
        sheet.conditional_formatting.add(
            f"{letter}2:{letter}{MAX_IMPORT_ROWS + 1}",
            FormulaRule(formula=[f'LEN(TRIM({letter}2))=0'], fill=red_fill),
        )


def build_template() -> bytes:
    workbook = Workbook()
    product_sheet = workbook.active
    product_sheet.title = PRODUCT_SHEET
    _style_sheet(product_sheet, PRODUCT_COLUMNS)
    _add_validations(product_sheet, product=True)

    variant_sheet = workbook.create_sheet(VARIANT_SHEET)
    _style_sheet(variant_sheet, VARIANT_COLUMNS)
    _add_validations(variant_sheet, product=False)

    example_sheet = workbook.create_sheet(EXAMPLE_SHEET)
    _style_sheet(example_sheet, PRODUCT_COLUMNS, example=True)
    example_sheet["A4"] = "SKU 示例（请填写到 SKU导入 工作表）"
    example_sheet["A4"].font = Font(bold=True, color="126B57")
    for index, (title, _key, _description, sample) in enumerate(VARIANT_COLUMNS, start=1):
        example_sheet.cell(5, index, title).font = Font(bold=True, color="FFFFFF")
        example_sheet.cell(5, index).fill = PatternFill("solid", fgColor="126B57")
        example_sheet.cell(6, index, sample).fill = PatternFill("solid", fgColor="FFF4D6")
        example_sheet.cell(6, index).alignment = Alignment(wrap_text=True, vertical="top")

    help_sheet = workbook.create_sheet(HELP_SHEET)
    help_sheet.append(["工作表", "字段", "是否必填", "填写规则", "示例"])
    for cell in help_sheet[1]:
        cell.fill = PatternFill("solid", fgColor="126B57")
        cell.font = Font(color="FFFFFF", bold=True)
    for sheet_name, columns in ((PRODUCT_SHEET, PRODUCT_COLUMNS), (VARIANT_SHEET, VARIANT_COLUMNS)):
        for title, _key, description, sample in columns:
            help_sheet.append([sheet_name, title.rstrip("*"), "是" if title.endswith("*") else "否", description, sample])
    help_sheet.freeze_panes = "A2"
    help_sheet.auto_filter.ref = f"A1:E{help_sheet.max_row}"
    for column, width in {"A": 14, "B": 20, "C": 12, "D": 54, "E": 42}.items():
        help_sheet.column_dimensions[column].width = width
    for row in help_sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "是" if value else "否"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _optional(value: Any) -> str | None:
    return _text(value) or None


def _split_items(value: Any) -> list[str]:
    return [item.strip() for item in re.split(r"[\n|；;]+", _text(value)) if item.strip()]


def _key_values(value: Any) -> dict[str, str]:
    result: dict[str, str] = {}
    for item in _split_items(value):
        separator = "=" if "=" in item else "：" if "：" in item else ":" if ":" in item else None
        if separator is None:
            raise ValueError(f"“{item}”缺少 = 分隔符")
        key, item_value = item.split(separator, 1)
        if not key.strip() or not item_value.strip():
            raise ValueError(f"“{item}”的名称或值为空")
        result[key.strip()] = item_value.strip()
    return result


def _datetime(value: Any) -> datetime | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.replace(tzinfo=value.tzinfo or timezone.utc)
    if isinstance(value, date):
        return datetime.combine(value, time.min, tzinfo=timezone.utc)
    text = _text(value).replace("/", "-")
    for pattern in ("%Y-%m-%d %H:%M", "%Y-%m-%d", "%Y-%m-%dT%H:%M"):
        try:
            return datetime.strptime(text, pattern).replace(tzinfo=timezone.utc)
        except ValueError:
            continue
    raise ValueError("时间格式应为 2026-09-18 10:30")


def _headers(sheet) -> dict[str, int]:
    return {_text(cell.value).rstrip("*"): index for index, cell in enumerate(sheet[1]) if _text(cell.value)}


def _row_values(sheet, row: int, columns: list[tuple[str, str, str, Any]]) -> dict[str, Any]:
    headers = _headers(sheet)
    return {
        key: sheet.cell(row, headers[title.rstrip("*")] + 1).value if title.rstrip("*") in headers else None
        for title, key, _description, _sample in columns
    }


def _validation_message(error: ValidationError) -> str:
    messages = []
    for item in error.errors():
        field = ".".join(str(part) for part in item["loc"])
        messages.append(f"{field}：{item['msg']}")
    return "；".join(messages)


def parse_catalog(content: bytes) -> ParsedCatalog:
    result = ParsedCatalog()
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception:
        result.errors.append(ImportErrorItem("文件", 0, "无法读取 Excel，请使用下载的 .xlsx 模板"))
        return result
    if PRODUCT_SHEET not in workbook.sheetnames:
        result.errors.append(ImportErrorItem("文件", 0, f"缺少“{PRODUCT_SHEET}”工作表"))
        return result

    product_sheet = workbook[PRODUCT_SHEET]
    required_product_headers = {"平台商品ID", "商品名称", "商品简介"}
    missing = required_product_headers - set(_headers(product_sheet))
    if missing:
        result.errors.append(ImportErrorItem(PRODUCT_SHEET, 1, f"缺少必填列：{'、'.join(sorted(missing))}"))
        return result
    if product_sheet.max_row - 1 > MAX_IMPORT_ROWS:
        result.errors.append(ImportErrorItem(PRODUCT_SHEET, 0, f"单次最多导入 {MAX_IMPORT_ROWS} 件商品"))
        return result

    seen_product_ids: set[str] = set()
    seen_internal_codes: dict[str, int] = {}
    for row in range(2, product_sheet.max_row + 1):
        raw = _row_values(product_sheet, row, PRODUCT_COLUMNS)
        if not any(_text(value) for value in raw.values()):
            continue
        try:
            product_id = _text(raw["id"])
            if product_id in seen_product_ids:
                raise ValueError("同一文件中平台商品 ID 重复")
            platform_text = _text(raw["platform"]).lower() or "拼多多"
            status_text = _text(raw["status"]).lower() or "草稿"
            if platform_text not in PLATFORM_MAP:
                raise ValueError("平台只能填写拼多多、京东、淘宝/天猫或其他")
            if status_text not in PRODUCT_STATUS_MAP:
                raise ValueError("资料状态只能填写草稿、已发布或已下线")
            internal_code = _optional(raw["internal_code"])
            if internal_code:
                if internal_code in seen_internal_codes:
                    raise ValueError(f"内部商品编码与第 {seen_internal_codes[internal_code]} 行重复")
                seen_internal_codes[internal_code] = row
            product = ProductCreate(
                id=product_id,
                internal_code=internal_code,
                platform=PLATFORM_MAP[platform_text],
                status=PRODUCT_STATUS_MAP[status_text],
                name=_text(raw["name"]),
                brand=_optional(raw["brand"]),
                model=_optional(raw["model"]),
                category_name=_optional(raw["category_name"]),
                summary=_text(raw["summary"]),
                selling_points=_split_items(raw["selling_points"]),
                specifications=_key_values(raw["specifications"]),
                usage=_optional(raw["usage"]),
                suitable_for=_optional(raw["suitable_for"]),
                warnings=_optional(raw["warnings"]),
                after_sales_limits=_optional(raw["after_sales_limits"]),
                owner_name=_optional(raw["owner_name"]),
                reviewed_by=_optional(raw["reviewed_by"]),
                effective_at=_datetime(raw["effective_at"]),
                published_at=_datetime(raw["published_at"]),
                compliance_notes=_optional(raw["compliance_notes"]),
            )
            result.products.append(product)
            seen_product_ids.add(product.id)
        except ValidationError as exc:
            result.errors.append(ImportErrorItem(PRODUCT_SHEET, row, _validation_message(exc)))
        except (ValueError, TypeError) as exc:
            result.errors.append(ImportErrorItem(PRODUCT_SHEET, row, str(exc)))

    if VARIANT_SHEET in workbook.sheetnames:
        variant_sheet = workbook[VARIANT_SHEET]
        required_variant_headers = {"平台商品ID", "SKU ID", "SKU名称"}
        missing = required_variant_headers - set(_headers(variant_sheet))
        if missing:
            result.errors.append(ImportErrorItem(VARIANT_SHEET, 1, f"缺少必填列：{'、'.join(sorted(missing))}"))
            return result
        if variant_sheet.max_row - 1 > MAX_IMPORT_ROWS:
            result.errors.append(ImportErrorItem(VARIANT_SHEET, 0, f"单次最多导入 {MAX_IMPORT_ROWS} 个 SKU"))
            return result
        seen_skus: set[str] = set()
        for row in range(2, variant_sheet.max_row + 1):
            raw = _row_values(variant_sheet, row, VARIANT_COLUMNS)
            if not any(_text(value) for value in raw.values()):
                continue
            try:
                sku_id = _text(raw["sku_id"])
                if sku_id in seen_skus:
                    raise ValueError("同一文件中 SKU ID 重复")
                stock_text = _text(raw["stock_status"]).lower() or "未知"
                status_text = _text(raw["status"]).lower() or "启用"
                if stock_text not in STOCK_STATUS_MAP:
                    raise ValueError("库存状态只能填写未知、有货、库存紧张或缺货")
                if status_text not in VARIANT_STATUS_MAP:
                    raise ValueError("SKU 状态只能填写启用或停用")
                quantity_text = _text(raw["stock_quantity"])
                price_text = _text(raw["price"])
                variant = VariantCreate(
                    sku_id=sku_id,
                    name=_text(raw["name"]),
                    attributes=_key_values(raw["attributes"]),
                    currency=_text(raw["currency"]) or "CNY",
                    price=Decimal(price_text) if price_text else None,
                    stock_status=STOCK_STATUS_MAP[stock_text],
                    stock_quantity=int(quantity_text) if quantity_text else None,
                    status=VARIANT_STATUS_MAP[status_text],
                )
                result.variants.append(VariantImport(row=row, product_id=_text(raw["product_id"]), variant=variant))
                seen_skus.add(variant.sku_id)
            except ValidationError as exc:
                result.errors.append(ImportErrorItem(VARIANT_SHEET, row, _validation_message(exc)))
            except (ValueError, TypeError) as exc:
                result.errors.append(ImportErrorItem(VARIANT_SHEET, row, str(exc)))
    if not result.products and not result.variants and not result.errors:
        result.errors.append(ImportErrorItem(PRODUCT_SHEET, 0, "没有找到可导入的数据，请从第 2 行开始填写"))
    return result
